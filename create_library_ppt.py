from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


OUT = "Library_Management_System_Overview.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(20, 37, 63)
BLUE = RGBColor(43, 102, 214)
TEAL = RGBColor(29, 163, 154)
GOLD = RGBColor(242, 176, 53)
INK = RGBColor(40, 52, 70)
MUTED = RGBColor(101, 116, 135)
PALE = RGBColor(243, 247, 252)
WHITE = RGBColor(255, 255, 255)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
            font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def bullet_box(slide, x, y, w, h, bullets, size=12, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def header(slide, kicker, title, subtitle):
    rect(slide, 0, 0, 13.333, 7.5, PALE)
    rect(slide, 0, 0, 13.333, 0.18, TEAL)
    textbox(slide, 0.68, 0.52, 3.0, 0.28, kicker.upper(), 10, TEAL, True)
    textbox(slide, 0.68, 0.84, 11.8, 0.55, title, 28, NAVY, True)
    textbox(slide, 0.7, 1.48, 11.8, 0.35, subtitle, 12, MUTED)


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "01  |  Member experience", "Everything members need, in a few clicks", "A simple self-service journey from registration to discovering and borrowing books.")
cards = [
    (0.7, "01", "Member registration & login", "Create a secure account and access a personalized member portal.", ["Quick sign-up with contact details", "Secure login and profile access"], BLUE),
    (4.55, "02", "Borrowing a book", "Search availability, check due dates, and keep borrowing history in one place.", ["Real-time availability status", "Clear due-date reminders"], TEAL),
    (8.4, "03", "Reserve a book", "Join a waitlist when a title is unavailable and get notified when it is ready.", ["One-click reservations", "Automated pickup notifications"], GOLD),
]
for x, num, title, desc, bullets, accent in cards:
    rect(slide, x, 2.25, 3.45, 4.35, WHITE, True, RGBColor(225, 232, 241))
    rect(slide, x, 2.25, 3.45, 0.13, accent)
    rect(slide, x+0.28, 2.65, 0.55, 0.55, accent, True)
    textbox(slide, x+0.28, 2.77, 0.55, 0.22, num, 10, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, x+0.28, 3.45, 2.85, 0.62, title, 17, NAVY, True)
    textbox(slide, x+0.28, 4.2, 2.78, 0.78, desc, 12, MUTED)
    bullet_box(slide, x+0.28, 5.25, 2.8, 0.9, bullets, 11)

textbox(slide, 0.7, 7.08, 11.8, 0.2, "Member portal  •  Search  •  Borrow  •  Reserve  •  Track", 10, MUTED, False, PP_ALIGN.CENTER)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "02  |  Librarian workspace", "A clear UI powered by data cloud insights", "Give librarians the operational view they need to manage circulation, members, and collection health.")

# Left dashboard mockup
rect(slide, 0.7, 2.18, 7.0, 4.55, NAVY, True)
textbox(slide, 1.0, 2.48, 2.4, 0.32, "LIBRARY ADMIN", 11, WHITE, True)
textbox(slide, 1.0, 2.88, 3.0, 0.36, "Good morning, librarian", 18, WHITE, True)
textbox(slide, 1.0, 3.23, 5.4, 0.25, "Your collection at a glance", 10, RGBColor(181, 198, 220))
stats = [(1.0, "1,248", "Active members", BLUE), (3.15, "326", "Books on loan", TEAL), (5.3, "42", "Reservations", GOLD)]
for x, val, label, color in stats:
    rect(slide, x, 3.75, 1.95, 1.03, RGBColor(35, 56, 86), True)
    textbox(slide, x+0.15, 3.93, 1.65, 0.3, val, 22, color, True, PP_ALIGN.CENTER)
    textbox(slide, x+0.15, 4.33, 1.65, 0.22, label, 9, RGBColor(199, 211, 227), False, PP_ALIGN.CENTER)
textbox(slide, 1.0, 5.2, 2.5, 0.25, "Borrowing activity", 11, WHITE, True)
# chart axes and bars
rect(slide, 1.0, 6.18, 5.7, 0.02, RGBColor(80, 103, 134))
for x, h, color in [(1.4, .48, BLUE), (2.05, .78, TEAL), (2.7, .62, BLUE), (3.35, 1.02, TEAL), (4.0, .7, BLUE), (4.65, 1.28, TEAL), (5.3, .94, BLUE), (5.95, 1.48, TEAL)]:
    rect(slide, x, 6.18-h, 0.38, h, color, True)
for i, m in enumerate(["May", "Jun", "Jul", "Aug"]):
    textbox(slide, 1.35+i*1.3, 6.28, 0.7, 0.18, m, 8, RGBColor(181, 198, 220), False, PP_ALIGN.CENTER)

# Right insights panel
rect(slide, 8.05, 2.18, 4.58, 4.55, WHITE, True, RGBColor(225, 232, 241))
textbox(slide, 8.38, 2.52, 3.7, 0.3, "Cloud insights", 17, NAVY, True)
textbox(slide, 8.38, 2.9, 3.65, 0.42, "Turn activity into action with live, centralized data.", 11, MUTED)
insights = [
    ("Collection health", "Identify popular and underused titles", TEAL),
    ("Member trends", "Spot demand patterns and peak periods", BLUE),
    ("Smart alerts", "Overdue, reservation, and inventory signals", GOLD),
]
for i, (t, d, c) in enumerate(insights):
    y = 3.65 + i*0.88
    rect(slide, 8.38, y, 0.12, 0.55, c, True)
    textbox(slide, 8.7, y+0.02, 3.45, 0.22, t, 11, NAVY, True)
    textbox(slide, 8.7, y+0.27, 3.45, 0.25, d, 10, MUTED)
textbox(slide, 8.38, 6.27, 3.7, 0.22, "One workspace. Better decisions.", 10, TEAL, True)

prs.save(OUT)
print(OUT)
